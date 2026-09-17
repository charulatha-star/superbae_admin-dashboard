"""Generates ADMIN_DASHBOARD_DOCUMENTATION.docx from the actual frontend implementation."""
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)

doc = Document()

# ---------- base styles ----------
for lvl, size in [(1, 16), (2, 13), (3, 11.5)]:
    st = doc.styles[f'Heading {lvl}']
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = PRIMARY
    st.font.bold = True
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)

def h1(text):
    p = doc.add_heading(text, level=1)
    p.paragraph_format.page_break_before = True
    return p

def h2(text): return doc.add_heading(text, level=2)
def h3(text): return doc.add_heading(text, level=3)

def para(text, italic=False, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.bold = bold
    if color: r.font.color.rgb = color
    return p

def bullets(items, style='List Bullet'):
    for it in items:
        if isinstance(it, tuple):  # (bold prefix, rest)
            p = doc.add_paragraph(style=style)
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            doc.add_paragraph(it, style=style)

def numbered(items):
    bullets(items, style='List Number')

def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Consolas'
    r.font.size = Pt(9)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    # light shading
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), 'F2F2F2')
    p.paragraph_format.element.get_or_add_pPr().append(shd)
    return p

def set_cell(cell, text, bold=False, header=False):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.font.size = Pt(9)
    r.bold = bold or header
    if header:
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear'); shd.set(qn('w:fill'), '1F4E79')
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell._tc.get_or_add_tcPr().append(shd)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        set_cell(t.rows[0].cells[i], hd, header=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell(cells[i], val)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return t

def note(text):
    para(text, italic=True, color=GREY)

def add_page_numbers():
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Admin Dashboard Documentation  |  Page ')
    r.font.size = Pt(8); r.font.color.rgb = GREY
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run('superbae_admin-dashboard  |  admin-dashboard-frontend')
    hr.font.size = Pt(8); hr.font.color.rgb = GREY

def add_toc():
    p = doc.add_paragraph()
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), r'TOC \o "1-2" \h \z \u')
    run = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = 'Right-click and choose "Update Field" to build the Table of Contents.'
    run.append(t)
    fld.append(run)
    p._p.append(fld)

# ---------- Title page ----------
for _ in range(6): doc.add_paragraph()
tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run('ADMIN DASHBOARD'); r.font.size = Pt(34); r.bold = True; r.font.color.rgb = PRIMARY
tp2 = doc.add_paragraph(); tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp2.add_run('Frontend Documentation'); r.font.size = Pt(20); r.font.color.rgb = GREY
doc.add_paragraph()
tp3 = doc.add_paragraph(); tp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp3.add_run('Project: admin-dashboard-frontend\nRepository: superbae_admin-dashboard\nBased on the current implementation as found in the repository\nVersion 1.0')
r.font.size = Pt(12); r.font.color.rgb = GREY
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

doc.add_heading('Table of Contents', level=1)
add_toc()

# ============================================================
# 1. Executive Summary
# ============================================================
h1('1. Executive Summary')
para('The Admin Dashboard is a web-based back-office application for administering a social / lifestyle mobile platform. '
     'It allows internal administrator accounts to manage end users, community content (posts, comments, reports, anonymous forum, partners, groups, events, trips), '
     'platform content (CMS: affirmations, zodiac, tips, banners, journal prompts, wardrobe), AI features, subscriptions/payments/revenue, support tickets, notifications, '
     'analytics, safety moderation, security events, audit logs, application settings, and other administrator accounts and their roles/permissions.')
para('The application is a Next.js (App Router) client application rendered mostly on the client ("use client"). '
     'It consumes an Express + MongoDB backend API (hosted in the sibling admin-dashboard-backend project) on port 3001 using JSON over HTTP with Bearer-token authentication.')
h2('1.1 Target Users')
bullets([
    'Super Admins (role permissions containing "*" wildcard) who have full access to all modules.',
    'Scoped administrators (e.g. role_operations is used as the default role id in code) who only see modules allowed by their role permissions.',
])
h2('1.2 Implementation Status Overview')
table(['Status', 'Modules / Pages'], [
    ['Implemented (API-backed)', 'Login, Dashboard, Users (list/edit/view), Admins (CRUD), Roles (CRUD), Community (posts, comments, reported posts/comments, moderation queue, blocked users, anonymous forum, partners), Groups, Events, Trips, Content CMS sub-pages, AI Dashboard/Usage/Config, Analytics Users, Finance sub-pages, Support Tickets, Notifications, Safety Reports/Moderation, Security, Audit Logs, App Settings, Trackers'],
    ['Placeholder pages (hardcoded static)', '/admin/ai, /admin/analytics, /admin/finance, /admin/safety index pages; /admin/analytics/features and /admin/analytics/revenue show hardcoded demo numbers'],
    ['Partially implemented / disabled', 'users/[id]/view page contains a fully commented-out earlier implementation above the active code; several Sidebar items and community counters are commented out (anonymous/groups counters, CMS sub-items, Settings/Profile footer links)'],
    ['Not determinable from frontend', 'Exact business rules for stats, subscription data semantics, permission catalog contents (fetched from /permissions but owned by backend)'],
])
note('All statements in this document refer to code in admin-dashboard-frontend unless explicitly noted otherwise.')

# ============================================================
# 2. Project Overview
# ============================================================
h1('2. Project Overview')
table(['Item', 'Value'], [
    ['Project name', 'admin-dashboard (package.json "name")'],
    ['Version', '0.1.0 (private)'],
    ['Framework', 'Next.js 16.3.2 (App Router), React 19.2.8, TypeScript 5'],
    ['Styling', 'Tailwind CSS 4 (@tailwindcss/postcss) + CSS Modules (*.module.css)'],
    ['UI / icons / charts', 'lucide-react icons, custom components, recharts 3 charts'],
    ['Forms & validation', 'react-hook-form 7 + zod 4 on the login page; manual state forms elsewhere'],
    ['State management', 'React local state + module-level shared singletons (no Redux/Zustand/React Query)'],
    ['HTTP client', 'Native fetch wrapped in fetchApi() (src/lib/api/api.ts)'],
    ['Backend API', 'Express 5 + Mongoose (admin-dashboard-backend), default base URL http://localhost:3001'],
    ['Mock tooling', 'json-server script "mock-api" (data/db.json) — db.json not present in repo'],
    ['Tests', 'No unit/e2e test runner configured in frontend package.json'],
])
h2('2.1 NPM Scripts (frontend)')
table(['Script', 'Command', 'Purpose'], [
    ['dev', 'next dev --hostname 0.0.0.0', 'Development server (LAN-accessible)'],
    ['build', 'next build', 'Production build'],
    ['start', 'next start', 'Serve production build'],
    ['lint', 'eslint', 'Linting'],
    ['dev:api', 'npm run dev --prefix backend', 'Start backend (expects ./backend folder)'],
    ['seed', 'npm run seed --prefix backend', 'Seed backend data'],
    ['mock-api', 'json-server --watch data/db.json --port 3001', 'Local mock API (data file absent)'],
])
h2('2.2 Related Backend (context only)')
para('The sibling folder admin-dashboard-backend is an Express + MongoDB API ("Express + MongoDB API replacing json-server" per its package.json) '
     'with seed and test:* scripts (test:auth, test:overview, test:security, test:profile, test:anonymous-moderation). '
     'It is referenced here only where it explains frontend behavior; its internals are out of scope.')

# ============================================================
# 3. Project Architecture
# ============================================================
h1('3. Project Architecture')
para('The frontend follows the Next.js App Router structure: each folder under src/app is a route; pages are client components that fetch data through a thin API layer (src/lib/api) using a shared fetchApi() wrapper. Authentication and permission state live in module-level singletons shared via custom hooks instead of a global store library.')
code('Admin (browser)\n'
     '   |\n'
     '   v\n'
     'Next.js App Router pages (src/app/**/page.tsx, "use client")\n'
     '   |\n'
     '   v\n'
     'Custom hooks (src/hooks: useAuth, usePermissions, useSidebar)\n'
     '   |                    |\n'
     '   v                    v\n'
     'UI components (src/components/admin/*)   Permission checks (PermissionGate)\n'
     '   |\n'
     '   v\n'
     'API services (src/lib/api: api.ts, auth.ts, admins.ts, roles.ts, users.ts)\n'
     '   |\n'
     '   v\n'
     'fetchApi() HTTP wrapper — Bearer token from sessionStorage,\n'
     '401 handling, { data,total,pages } envelope normalization\n'
     '   |\n'
     '   v\n'
     'Backend API (Express, http://localhost:3001)  -->  MongoDB (Mongoose)')
h2('3.1 Key Architectural Characteristics')
bullets([
    ('Client-side rendering: ', 'nearly all pages are marked "use client"; data is fetched after mount with skeleton loaders during requests.'),
    ('Central HTTP wrapper: ', 'all requests go through fetchApi() in src/lib/api/api.ts, which injects the Bearer token, handles 401 by redirecting to /login, extracts server error messages, and unwraps { data, total, pages } envelopes.'),
    ('Shared auth singleton: ', 'useAuth() hoists admin state to module level so only one /auth/me call fires per page load (src/hooks/useAuth.ts).'),
    ('Role caching: ', 'usePermissions() caches the admin role in sessionStorage (role_<roleId>) to avoid duplicate /roles/:id calls (src/hooks/usePermissions.ts).'),
    ('Permission-gated UI: ', 'PermissionGate component and Sidebar filtering hide actions/routes the role lacks.'),
    ('No global store: ', 'no Redux, Zustand, React Context providers, or server-state libraries are used.'),
])

# ============================================================
# 4. Project Folder Structure
# ============================================================
h1('4. Project Folder Structure')
table(['Path', 'Responsibility'], [
    ['src/app/', 'App Router routes: /login, /403, /admin/** modules, root layout, not-found'],
    ['src/app/admin/layout.tsx', 'Client layout guarding /admin via useAuth(); renders AdminLayout (Navbar, Sidebar, Footer)'],
    ['src/app/admin/<module>/', 'One folder per admin module: users, admins, roles, community, content, ai, analytics, finance, safety, settings, support, trips, groups, events, notifications, profile, security, audit-logs, dashboard'],
    ['src/components/admin/', 'AdminLayout, Navbar, Sidebar, Footer, Loader, Skeleton, Toast, ConfirmModal, EmptyState, ErrorState, NoData, PasswordInput, Tabs, Accordion, PermissionGate'],
    ['src/hooks/', 'useAuth (shared auth state, redirect, logout), usePermissions (role fetch/cache, hasPermission), useSidebar (collapse state persisted in localStorage)'],
    ['src/lib/api/', 'api.ts (fetchApi wrapper + token storage), auth.ts, admins.ts, roles.ts, users.ts (typed service functions)'],
    ['src/lib/mongo/mongoose.ts', 'Server-side mongoose connect helper (MONGODB_URI, fallback mongodb://127.0.0.1:27017/admin-dashboard). Not referenced by app pages — likely legacy of server-side data fetching'],
    ['src/types/', 'admin.ts (Admin), role.ts (Role, Permission)'],
    ['docs_gen/', 'This documentation generator (added by documentation task)'],
])

# ============================================================
# 5. Admin Panel Features (modules)
# ============================================================
h1('5. Admin Panel Features')
table(['Module', 'Purpose', 'Status'], [
    ['Dashboard', 'Role-specific stat cards, activity line chart (recharts) and recent-activity table driven by /dashboardStats, /dashboardCharts, /dashboardTables keyed by admin.roleId', 'Implemented'],
    ['Users', 'End-user list with search/filter/sort/pagination, stats, inline edit modal, delete confirm modal; detail view with tabbed data (overview, profile, security, subscription, events, community, personal, activity), block/unblock and status actions', 'Implemented'],
    ['Admins', 'CRUD for administrator accounts: list + create form + edit page, delete confirmation, role assignment', 'Implemented'],
    ['Roles & Permissions', 'CRUD for roles; UI blocks deleting roles that have assigned admins; permissions catalog fetched from /permissions', 'Implemented'],
    ['Community', 'Hub page with counters; posts list + detail (delete post/comment), comments list (delete), reported posts/comments (status updates), moderation queue merging reportedPosts/reportedComments/safetyModeration, blocked users (unblock), partners list', 'Implemented'],
    ['Anonymous forum moderation', 'Anonymous posts queue with filter/search/pagination, post detail, moderation action (hide/remove/warn/suspend/ban/restore) with reason, paginated moderation history', 'Implemented'],
    ['Groups / Events / Trips', 'Read-only lists from /groups, /events, /trips', 'Implemented (read-only)'],
    ['Content (CMS)', 'Hub page; sub-pages list affirmations, banners, journal prompts, tips, wardrobe, zodiac', 'Implemented (read-only)'],
    ['AI', 'AI Dashboard metrics (/aiDashboard), usage list (/aiUsage), configuration list (/aiConfig)', 'Implemented (read-only)'],
    ['Analytics', 'User analytics table (/analyticsUsers); feature & revenue analytics pages show hardcoded demo values', 'Partially implemented'],
    ['Finance', 'Payments, revenue and subscriptions read-only lists', 'Implemented (read-only)'],
    ['Support / Notifications / Security / Audit Logs', 'Read-only lists (/support, /notifications, /security, /auditLogs)', 'Implemented (read-only)'],
    ['Safety', 'Reports and moderation queues (/safetyReports, /safetyModeration); module index page is a placeholder', 'Implemented (sub-pages)'],
    ['Settings', 'Read-only: /settings, /appSettings, /trackers', 'Implemented (read-only)'],
    ['Profile', 'Profile form from session admin (name, email, password, avatar); no persistence API call found', 'Partially implemented'],
])

# ============================================================
# 6. Routes and Pages
# ============================================================
h1('6. Routes and Pages')
para('All routes are file-system routes of the Next.js App Router. Routes under /admin are guarded by src/app/admin/layout.tsx (via useAuth redirect). '
     'The sidebar (src/components/admin/Sidebar/Sidebar.tsx) additionally hides links the current role lacks permission for.')
table(['Route', 'Page', 'Purpose', 'Primary APIs', 'Auth', 'Notes'], [
    ['/', 'RootPage', 'Redirects to /login', '—', 'None', 'src/app/page.tsx'],
    ['/login', 'LoginPage', 'Email/password login (react-hook-form + zod)', 'POST /auth/login', 'None', 'Stores admin + token in sessionStorage'],
    ['/403', 'ForbiddenPage', 'Access-denied screen', '—', 'None', 'src/app/403/page.tsx'],
    ['(unmatched)', 'NotFoundPage', '404 screen', '—', 'None', 'src/app/not-found.tsx'],
    ['/admin/dashboard', 'DashboardPage', 'Stats, chart, recent activity (role-keyed)', 'GET /dashboardStats, /dashboardCharts, /dashboardTables', 'Required', 'recharts LineChart'],
    ['/admin/users', 'UsersPage', 'User list: search, filter, sort, pagination (10/page), edit modal, delete modal', 'GET /users, PATCH /users/:id, DELETE /users/:id', 'Required; users.view/edit', ''],
    ['/admin/users/[id]', 'User edit', 'Edit user fields', 'GET /users/:id, /users/:id/profile, /personal, PATCH /users/:id, GET /sessions?userId=, /loginHistory?userId=', 'Required; users.edit', ''],
    ['/admin/users/[id]/view', 'User detail (tabs)', 'Overview/Profile/Events/Community/Personal/Security/Subscription/Activity tabs; account actions', 'GET /users/:id + tab endpoints (with ?page=), PATCH /users/:id', 'Required; users.edit', 'Has commented-out earlier implementation'],
    ['/admin/admins', 'AdminsPage', 'Admin accounts list, delete, pagination', 'GET /admins, GET /roles, DELETE /admins/:id', 'Required; admins.* gates', ''],
    ['/admin/admins/create', 'Create admin', 'Create administrator', 'POST /admins', 'Required', ''],
    ['/admin/admins/[id]', 'Admin edit', 'Edit administrator', 'GET /admins/:id, PATCH /admins/:id', 'Required', ''],
    ['/admin/roles', 'RolesPage', 'Roles list, admin-count, delete guard', 'GET /roles, GET /admins, DELETE /roles/:id', 'Required; roles.* gates', ''],
    ['/admin/roles/create', 'Create role', 'Create role', 'POST /roles', 'Required', ''],
    ['/admin/roles/[id]', 'Role edit', 'Edit role', 'GET /roles/:id, PATCH /roles/:id', 'Required', ''],
    ['/admin/community', 'CommunityPage', 'Hub + counters', 'GET /posts, /comments, /reportedPosts, /reportedComments, /safetyModeration, /users', 'Required; posts.view', 'Anonymous/groups counters commented out'],
    ['/admin/community/posts', 'Posts list', 'List/delete posts', 'GET /posts, DELETE /posts/:id', 'Required; community.posts.view', ''],
    ['/admin/community/posts/[id]', 'Post detail', 'Post + comments, delete comment', 'GET /posts/:id, GET /comments, DELETE /comments/:id', 'Required', ''],
    ['/admin/community/comments', 'Comments list', 'List/delete comments', 'GET /comments, DELETE /comments/:id', 'Required', ''],
    ['/admin/community/reported/posts', 'Reported posts', 'Reports with status updates', 'GET /reportedPosts, PATCH /reportedPosts/:id', 'Required', 'Shared ReportedList.tsx'],
    ['/admin/community/reported/comments', 'Reported comments', 'Reports with status updates', 'GET /reportedComments, PATCH /reportedComments/:id', 'Required', 'Shared ReportedList.tsx'],
    ['/admin/community/moderation', 'Moderation queue', 'Merged queue with actions', 'GET /reportedPosts, /reportedComments, /safetyModeration, PATCH /<resource>/<id>', 'Required', ''],
])

table(['Route', 'Page', 'Purpose', 'Primary APIs', 'Auth', 'Notes'], [
    ['/admin/community/blocked', 'Blocked users', 'List blocked users, unblock', 'GET /users (client filter), PATCH /users/:id', 'Required', ''],
    ['/admin/community/anonymous', 'Anonymous queue', 'Queue with filter/search/pagination', 'GET /anonymousPosts?filter&search&page&limit', 'Required; posts.view', 'preserveEnvelope=True'],
    ['/admin/community/anonymous/[postId]/view', 'Anonymous detail', 'Detail + moderation action', 'GET /anonymousPosts/:postId, POST /anonymousPosts/:postId/action', 'Required', ''],
    ['/admin/community/anonymous/history', 'History', 'Audit trail with filters', 'GET /anonymousModerationHistory?page&limit', 'Required', ''],
    ['/admin/community/partners', 'Partners', 'Partners list', 'GET /partners', 'Required', ''],
    ['/admin/groups', 'Groups', 'Groups list', 'GET /groups', 'Required; groups.view', ''],
    ['/admin/events', 'Events', 'Events list', 'GET /events', 'Required; events.view', ''],
    ['/admin/trips', 'Trips', 'Trips list', 'GET /trips', 'Required; events.view', ''],
    ['/admin/content', 'CMS hub', 'Links to content pages', '—', 'Required; cms.view', 'Static hub'],
    ['/admin/content/(affirmations|banners|journal|tips|wardrobe|zodiac)', 'CMS lists', 'Read-only lists', 'GET /affirmations, /banners, /journalPrompts, /tips, /wardrobe, /zodiac', 'Required', ''],
    ['/admin/ai', 'AI index', 'Placeholder', '—', '—', 'Hardcoded placeholder'],
    ['/admin/ai/dashboard | usage | config', 'AI pages', 'AI metrics / usage / config', 'GET /aiDashboard, /aiUsage, /aiConfig', 'Required; ai.view / ai.configure', ''],
    ['/admin/analytics', 'Analytics index', 'Placeholder', '—', '—', 'Hardcoded placeholder'],
    ['/admin/analytics/users', 'User analytics', 'Analytics table', 'GET /analyticsUsers', 'Required; analytics.users', ''],
    ['/admin/analytics/features | revenue', 'Static analytics', 'Static demo stats', '—', '—', 'Hardcoded values'],
    ['/admin/audit-logs', 'Audit logs', 'Log list', 'GET /auditLogs', 'Required; audit_logs.view', ''],
    ['/admin/finance', 'Finance index', 'Placeholder', '—', '—', 'Hardcoded placeholder'],
    ['/admin/finance/(payments|revenue|subscriptions)', 'Finance lists', 'Read-only lists', 'GET /payments, /revenue, /subscriptions', 'Required; payments.view / revenue.view / subscriptions.view', ''],
    ['/admin/notifications', 'Notifications', 'Notification list', 'GET /notifications', 'Required; users.view', ''],
    ['/admin/profile', 'ProfilePage', 'Profile form from session admin', 'None found', 'Required', 'No persistence API call found'],
    ['/admin/safety', 'Safety index', 'Placeholder', '—', '—', 'Hardcoded placeholder'],
    ['/admin/safety/reports | moderation', 'Safety pages', 'Reports / moderation items', 'GET /safetyReports, /safetyModeration', 'Required; reports.view / reports.resolve', ''],
    ['/admin/security', 'Security', 'Security events', 'GET /security', 'Required; security.view', ''],
    ['/admin/settings', 'Settings', 'Settings object', 'GET /settings', 'Required', ''],
    ['/admin/settings/app | trackers', 'Settings pages', 'App settings / trackers', 'GET /appSettings, /trackers', 'Required; security.view', ''],
    ['/admin/support', 'Support tickets', 'Ticket list', 'GET /support', 'Required; support.view', ''],
])
note('Permission strings are exactly as used in code: users.view, users.edit, admins.create/edit/delete/view, roles.create/edit/delete/view, community.posts.view, groups.view, events.view, cms.view, ai.view, ai.configure, subscriptions.view, payments.view, revenue.view, support.view, analytics.users, analytics.features, analytics.revenue, reports.view, reports.resolve, audit_logs.view, security.view.')

# ============================================================
# 7. Authentication
# ============================================================
h1('7. Authentication')
h2('7.1 Login Flow')
numbered([
    'Admin submits email/password on /login (src/app/login/page.tsx). The form is validated by a zod schema: valid email, password required (min length 1).',
    'loginAdmin(email, password) (src/lib/api/auth.ts) sends POST /auth/login with JSON body { email, password }.',
    'On success the response is an Admin object (src/types/admin.ts) which may include a token field.',
    'The page stores the full admin JSON under sessionStorage key "admin" and, if admin.token exists, stores the token via setAuthToken() (TOKEN_KEY is also "admin" in src/lib/api/api.ts — note: the token and the admin JSON share the same key; the token write happens last).',
    'The user is redirected to /admin/dashboard.',
    'On failure, the server error message (body.message) or "Invalid email or password." is shown in an error alert.',
])
h2('7.2 Session Handling')
table(['Aspect', 'Implementation'], [
    ['Token storage', 'sessionStorage key "admin" (TOKEN_KEY in src/lib/api/api.ts). Cleared on logout and on 401.'],
    ['Admin profile storage', 'sessionStorage key "admin" also holds the full Admin JSON written by the login page.'],
    ['Auth header', 'fetchApi() adds Authorization: Bearer <token> when a token is present.'],
    ['Session validation', 'GET /auth/me via getCurrentAdmin() (src/lib/api/auth.ts); errors are caught and return null.'],
    ['Shared auth state', 'useAuth() (src/hooks/useAuth.ts) hoists admin state to module scope with a single-flight fetch; subscribers share one /auth/me call per page load.'],
    ['Route protection', 'src/app/admin/layout.tsx renders Loader while loading and null when unauthenticated; useAuth() redirects unauthenticated users from /admin/* to /login and logged-in users from /login to /admin/dashboard.'],
    ['Logout', 'Sidebar logout calls logout() (useAuth.ts): POST /auth/logout first; only on success are token/session cleared and the user redirected to /login. If the API call fails, the session is kept.'],
    ['Expiry handling', 'A 401 from any API call clears the token/session and hard-redirects to /login with error "Session expired. Please log in again." (src/lib/api/api.ts). No refresh-token mechanism exists.'],
])
code('LoginPage (onSubmit)\n'
     '  -> loginAdmin()  [src/lib/api/auth.ts]\n'
     '  -> fetchApi() POST /auth/login\n'
     '  -> sessionStorage["admin"] = admin JSON; setAuthToken(admin.token)\n'
     '  -> router.push("/admin/dashboard")\n'
     '  -> /admin/layout.tsx useAuth() -> GET /auth/me (single shared call)\n'
     '  -> AdminLayout renders children')

# ============================================================
# 8. Authorization and Permissions
# ============================================================
h1('8. Authorization and Permissions')
h2('8.1 Model (as implemented)')
bullets([
    'Admin.roleId links an administrator to a Role (src/types/admin.ts, src/types/role.ts).',
    'Role.permissions is an array of permission strings. The wildcard "*" grants everything (Super Admin in usePermissions.hasPermission).',
    'usePermissions() (src/hooks/usePermissions.ts) fetches the role with GET /roles/:roleId, caches it in memory and sessionStorage (key role_<roleId>), and exposes hasPermission(permission).',
    'PermissionGate (src/components/admin/PermissionGate.tsx) renders children only when hasPermission(permission) is true; optional fallback otherwise; renders nothing while loading.',
    'Sidebar (src/components/admin/Sidebar/Sidebar.tsx) filters every nav item by its permission; items with permission "*" are always visible; empty sections are hidden.',
])
h2('8.2 Permission Checks Found in the Code')
table(['Location', 'Permission', 'Effect'], [
    ['Sidebar Users / Community / Anonymous', 'users.view / posts.view', 'Hide links'],
    ['Sidebar Groups / Events / Trips / Partners / CMS / Wardrobe', 'groups.view / events.view / users.view / cms.view', 'Hide links'],
    ['Sidebar AI section', 'ai.view / ai.configure', 'Hide links'],
    ['Sidebar Finance section', 'subscriptions.view / payments.view / revenue.view', 'Hide links'],
    ['Sidebar Support / Notifications', 'support.view / users.view', 'Hide links'],
    ['Sidebar Analytics section', 'analytics.users / analytics.features / analytics.revenue', 'Hide links'],
    ['Sidebar Safety section', 'reports.view / reports.resolve', 'Hide links'],
    ['Sidebar Administration section', 'admins.view / roles.view / audit_logs.view / security.view', 'Hide links'],
    ['Admins page buttons', 'admins.create / admins.edit / admins.delete', 'Hide create link, edit link, delete button per row'],
    ['Roles page buttons', 'roles.create / roles.edit / roles.delete', 'Hide create link, edit link, delete button per row'],
    ['Users page row actions', 'users.view / users.edit', 'Hide view link, edit link per row'],
    ['User edit page (whole page)', 'users.edit', 'Render nothing without permission'],
    ['User detail view', 'users.edit', 'Hide "Edit User" link'],
    ['Community posts list', 'community.posts.view', 'Hide view and delete actions per row'],
])
para('There is no client-side route-level permission guard: page protection is authentication-based (admin layout); permission checks are applied at component/UI level. '
     'Server-side enforcement cannot be verified from the frontend.', italic=True)
h2('8.3 Known Roles')
para('Role id referenced in code: role_operations (default roleId in registerAdmin, src/lib/api/auth.ts). '
     'Other role ids are dynamic (role_<timestamp> from createRole/createAdmin) or backend-seeded; dashboard stats are looked up by admin.roleId. '
     'The full role catalog lives in the backend — not determinable from the frontend implementation.')

# ============================================================
# 9. API Documentation (Authentication + Admins + Roles)
# ============================================================
h1('9. API Documentation')
para('Base URL: http://localhost:3001 in the browser (hardcoded in getApiBaseUrl(), src/lib/api/api.ts). '
     'Server-side: process.env.NEXT_PUBLIC_API_URL or http://localhost:3001. '
     'All requests set Content-Type: application/json and add Authorization: Bearer <token> when a token exists (login precedes token issuance). '
     'Non-2xx responses throw an Error using body.message when present, else "API error: <statusText>". '
     'A 401 clears the session and hard-redirects to /login. 204 responses yield undefined.')
note('Response field details are only documented where the shape is explicitly consumed in the code. '
     'For endpoints whose response content is treated as opaque (e.g. UserTabData = Record<string,unknown>, or raw arrays), fields are not determinable from the frontend implementation.')

h2('9.1 Authentication APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['POST', '/auth/login', 'Authenticate admin; returns Admin (with token)', '/login page', 'loginAdmin() — src/lib/api/auth.ts'],
    ['GET', '/auth/me', 'Fetch current admin session', 'useAuth() on every admin page', 'getCurrentAdmin() — src/lib/api/auth.ts'],
    ['POST', '/auth/logout', 'End session', 'Sidebar logout via useAuth().logout()', 'logoutAdmin() — src/lib/api/auth.ts'],
])
h3('POST /auth/login')
bullets([
    'Request body: { email: string, password: string }.',
    'Headers: Content-Type: application/json (no token before login).',
    'Response: Admin object; login page reads admin.token to persist the session.',
    'Errors: body.message or statusText shown in login error alert.',
])
h3('GET /auth/me')
bullets([
    'Requires Bearer token.',
    'Response: Admin object or null; getCurrentAdmin() catches errors, console.error logs, returns null.',
])
h3('POST /auth/logout')
bullets([
    'Requires Bearer token; request body not determinable from the frontend implementation.',
    'On success: token/admin cleared, redirect to /login. On failure: session intentionally kept (catch in logout).',
])

h2('9.2 Admins APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['GET', '/admins', 'List admins', '/admin/admins, /admin/roles', 'getAdmins() — src/lib/api/admins.ts'],
    ['GET', '/admins/:id', 'Get admin', '/admin/admins/[id]', 'getAdmin()'],
    ['POST', '/admins', 'Create admin', '/admin/admins/create; registerAdmin() in auth.ts', 'createAdmin()'],
    ['PATCH', '/admins/:id', 'Partial update admin', '/admin/admins/[id]', 'updateAdmin()'],
    ['DELETE', '/admins/:id', 'Delete admin', '/admin/admins (confirm modal)', 'deleteAdmin()'],
])
bullets([
    'POST createAdmin/request body: full Admin object built client-side, id = admin_<Date.now()>, createdAt = ISO timestamp.',
    'PATCH request body: Partial<Admin>.',
    'registerAdmin() (src/lib/api/auth.ts) POSTs /admins with defaults roleId=role_operations, status=active, twoFactorEnabled=false — not used by any UI.',
])

h2('9.3 Roles & Permissions APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['GET', '/roles', 'List roles', '/admin/roles, /admin/admins', 'getRoles() — src/lib/api/roles.ts'],
    ['GET', '/roles/:id', 'Get role (permissions)', 'usePermissions() on gated pages', 'getRole()'],
    ['POST', '/roles', 'Create role', '/admin/roles/create', 'createRole() (id = role_<Date.now()>)'],
    ['PATCH', '/roles/:id', 'Update role', '/admin/roles/[id]', 'updateRole()'],
    ['DELETE', '/roles/:id', 'Delete role', '/admin/roles (confirm modal)', 'deleteRole()'],
    ['GET', '/permissions', 'Permissions catalog', 'Defined in service; no page imports getPermissions()', 'getPermissions()'],
])
bullets([
    'Role object: { id, name, description, system?, permissions: string[] } (src/types/role.ts).',
    'Permission object: { id, module, resource, action } (src/types/role.ts); not currently consumed by any page despite getPermissions() existing.',
])

h2('9.4 Users APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['GET', '/users', 'List users', '/admin/users; /admin/community (counters, blocked list)', 'getUsers() — src/lib/api/users.ts'],
    ['GET', '/users/:id', 'Get user', '/admin/users/[id], /admin/users/[id]/view', 'getUser()'],
    ['POST', '/users', 'Create user', 'Service defined; no UI usage found', 'createUser()'],
    ['PUT', '/users/:id', 'Replace user', 'Service defined; no UI usage found', 'replaceUser()'],
    ['PATCH', '/users/:id', 'Partial update (edit modal, block/unblock, status actions)', 'users page edit modal; community/blocked unblock; users/[id] & users/[id]/view', 'updateUser()'],
    ['DELETE', '/users/:id', 'Delete user', 'users page delete modal', 'deleteUser()'],
    ['GET', '/users/:id/overview', 'Overview tab', '/admin/users/[id]/view', 'getUserOverview()'],
    ['GET', '/users/:id/security', 'Security tab data', '/admin/users/[id]/view, users/[id]', 'getUserSecurity()'],
    ['GET', '/users/:id/profile', 'Profile tab data', '/admin/users/[id]/view, users/[id]', 'getUserProfile()'],
    ['GET', '/users/:id/subscription', 'Subscription tab data', '/admin/users/[id]/view', 'getUserSubscription()'],
    ['GET', '/users/:id/content', 'Content tab data', 'Service defined; no UI usage found', 'getUserContent()'],
    ['GET', '/users/:id/events', 'Events tab (with ?page=)', '/admin/users/[id]/view', 'getUserEvents()'],
    ['GET', '/users/:id/community', 'Community tab (with ?page=)', '/admin/users/[id]/view', 'getUserCommunity()'],
    ['GET', '/users/:id/personal', 'Personal tab (with ?page=)', '/admin/users/[id]/view, users/[id]', 'getUserPersonal()'],
    ['GET', '/users/:id/activity', 'Activity tab (with ?page=)', '/admin/users/[id]/view', 'getUserActivity()'],
    ['GET', '/sessions?userId=:id', 'Device sessions', '/admin/users/[id]', 'direct fetchApi in page'],
    ['GET', '/loginHistory?userId=:id', 'Login history', '/admin/users/[id]', 'direct fetchApi in page'],
])
bullets([
    'GET /users/:id/activity returns sections: allActivity, timeline, goalActivity, journalActivity, trackerActivity, featureUsage (each an ActivitySectionData with { data?, total?, pages? }) per src/lib/api/users.ts.',
    'GET /users/:id/community returns an object with pagination info (communityPage state, .pages) — response consumed as CommunityData (not exported; not determinable precisely).'
])

h2('9.5 Community APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['GET', '/posts', 'List posts', '/admin/community (counter), posts list, post detail', 'fetchApi direct'],
    ['GET', '/posts/:id', 'Get post', '/admin/community/posts/[id]', 'fetchApi direct'],
    ['DELETE', '/posts/:id', 'Delete post', '/admin/community/posts', 'fetchApi direct'],
    ['GET', '/comments', 'List comments', '/admin/community (counter), post detail, comments list', 'fetchApi direct'],
    ['DELETE', '/comments/:id', 'Delete comment', '/admin/community/posts/[id], comments list', 'fetchApi direct'],
    ['GET', '/reportedPosts', 'List reported posts', '/admin/community (counter), moderation queue, reported posts list', 'fetchApi direct'],
    ['GET', '/reportedComments', 'List reported comments', '/admin/community (counter), moderation queue, reported comments list', 'fetchApi direct'],
    ['PATCH', '/reportedPosts/:id', 'Update report status', '/admin/community/reported/posts (ReportedList)', 'fetchApi via ReportedList'],
    ['PATCH', '/reportedComments/:id', 'Update report status', '/admin/community/reported/comments (ReportedList)', 'fetchApi via ReportedList'],
    ['GET', '/safetyModeration', 'Moderation queue items', '/admin/community (counter), moderation queue, safety/moderation', 'fetchApi direct'],
    ['GET', '/partners', 'List partners', '/admin/community/partners', 'fetchApi direct'],
    ['GET', '/users', 'Get users (filter blocked client-side)', '/admin/community/blocked', 'fetchApi direct'],
    ['PATCH', '/users/:id', 'Unblock user', '/admin/community/blocked', 'fetchApi direct'],
])
bullets([
    'ReportedList.tsx (src/app/admin/community/ReportedList.tsx) is a shared sub-component used by the reported-posts and reported-comments pages; it PATCHes /<resource>/<id> to update a report status.',
    'Anonymous endpoints below also support preserveEnvelope (raw { data, page, pages }).',
])

h2('9.6 Anonymous Forum Moderation APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['GET', '/anonymousPosts?filter=&search=&page=&limit=', 'Paginated anonymous posts queue', '/admin/community/anonymous', 'fetchApi direct (preserveEnvelope=true)'],
    ['GET', '/anonymousPosts/:postId', 'Get anonymous post detail', '/admin/community/anonymous/[postId]/view', 'fetchApi direct'],
    ['POST', '/anonymousPosts/:postId/action', 'Apply moderation action', '/admin/community/anonymous/[postId]/view', 'fetchApi direct'],
    ['GET', '/anonymousModerationHistory?page=&limit=', 'Paginated moderation history', '/admin/community/anonymous/history', 'fetchApi direct (preserveEnvelope=true)'],
])
bullets([
    'POST /anonymousPosts/:postId/action body: { action: string, reason?: string }; actions observed: hide, remove, warn, suspend, ban, restore (history filter options).',
    'Anonymous queue response is consumed as a raw envelope { data, page, pages } (preserveEnvelope=true) per ReportedList/queue pages.'
])

h2('9.7 Dashboard, Analytics, Finance, AI, Support, Safety, Security, Settings, Audit APIs')
table(['Method', 'Endpoint', 'Purpose', 'Used by', 'Service'], [
    ['GET', '/dashboardStats', 'Role-keyed dashboard stats', '/admin/dashboard', 'fetchApi direct'],
    ['GET', '/dashboardCharts', 'Role-keyed chart data', '/admin/dashboard', 'fetchApi direct'],
    ['GET', '/dashboardTables', 'Role-keyed recent activity table', '/admin/dashboard', 'fetchApi direct'],
    ['GET', '/analyticsUsers', 'User analytics rows', '/admin/analytics/users', 'fetchApi direct'],
    ['GET', '/payments', 'Payments list', '/admin/finance/payments', 'fetchApi direct'],
    ['GET', '/revenue', 'Revenue list', '/admin/finance/revenue', 'fetchApi direct'],
    ['GET', '/subscriptions', 'Subscriptions list', '/admin/finance/subscriptions', 'fetchApi direct (guarded Array.isArray)'],
    ['GET', '/support', 'Support tickets', '/admin/support', 'fetchApi direct'],
    ['GET', '/notifications', 'Notifications', '/admin/notifications', 'fetchApi direct'],
    ['GET', '/security', 'Security events', '/admin/security', 'fetchApi direct'],
    ['GET', '/auditLogs', 'Audit log entries', '/admin/audit-logs', 'fetchApi direct'],
    ['GET', '/appSettings', 'Application settings', '/admin/settings/app', 'fetchApi direct'],
    ['GET', '/trackers', 'Trackers', '/admin/settings/trackers', 'fetchApi direct'],
    ['GET', '/settings', 'Settings object', '/admin/settings', 'fetchApi direct'],
    ['GET', '/aiDashboard', 'AI dashboard metrics', '/admin/ai/dashboard', 'fetchApi direct'],
    ['GET', '/aiUsage', 'AI usage rows', '/admin/ai/usage', 'fetchApi direct'],
    ['GET', '/aiConfig', 'AI configuration rows', '/admin/ai/config', 'fetchApi direct'],
    ['GET', '/safetyReports', 'Safety reports', '/admin/safety/reports', 'fetchApi direct'],
    ['GET', '/groups', 'Groups list', '/admin/groups', 'fetchApi direct'],
    ['GET', '/events', 'Events list', '/admin/events', 'fetchApi direct'],
    ['GET', '/trips', 'Trips list', '/admin/trips', 'fetchApi direct'],
    ['GET', '/affirmations', 'Affirmations content', '/admin/content/affirmations', 'fetchApi direct'],
    ['GET', '/banners', 'Banners content', '/admin/content/banners', 'fetchApi direct'],
    ['GET', '/journalPrompts', 'Journal prompts content', '/admin/content/journal', 'fetchApi direct'],
    ['GET', '/tips', 'Wellness tips content', '/admin/content/tips', 'fetchApi direct'],
    ['GET', '/wardrobe', 'Wardrobe items content', '/admin/content/wardrobe', 'fetchApi direct'],
    ['GET', '/zodiac', 'Zodiac content', '/admin/content/zodiac', 'fetchApi direct'],
])
note('The analytics/features and analytics/revenue pages are STATIC (no fetchApi calls — hardcoded demo values). '
     'Content sub-pages and many management pages are read-only list displays consuming GET endpoints only.')

# ============================================================
# 10. API Inventory
# ============================================================
h1('10. API Inventory')
table(['Method', 'Endpoint', 'Purpose', 'Used By'], [
    ['POST', '/auth/login', 'Admin login', '/login'],
    ['GET', '/auth/me', 'Current admin', 'useAuth()'],
    ['POST', '/auth/logout', 'Logout', 'useAuth().logout()'],
    ['GET', '/admins', 'List admins', '/admin/admins, /admin/roles'],
    ['GET', '/admins/:id', 'Get admin', '/admin/admins/[id]'],
    ['POST', '/admins', 'Create admin', '/admin/admins/create'],
    ['PATCH', '/admins/:id', 'Update admin', '/admin/admins/[id]'],
    ['DELETE', '/admins/:id', 'Delete admin', '/admin/admins'],
    ['GET', '/roles', 'List roles', '/admin/roles, /admin/admins'],
    ['GET', '/roles/:id', 'Role + permissions', 'usePermissions()'],
    ['POST', '/roles', 'Create role', '/admin/roles/create'],
    ['PATCH', '/roles/:id', 'Update role', '/admin/roles/[id]'],
    ['DELETE', '/roles/:id', 'Delete role', '/admin/roles'],
    ['GET', '/permissions', 'Permission catalog', 'service only'],
    ['GET', '/users', 'List users', '/admin/users, /admin/community'],
    ['GET', '/users/:id', 'Get user', '/admin/users/[id], users/[id]/view'],
    ['POST', '/users', 'Create user', 'service only'],
    ['PUT', '/users/:id', 'Replace user', 'service only'],
    ['PATCH', '/users/:id', 'Update user', '/admin/users, users/[id], blocked'],
    ['DELETE', '/users/:id', 'Delete user', '/admin/users'],
    ['GET', '/users/:id/overview', 'Overview', 'users/[id]/view'],
    ['GET', '/users/:id/security', 'Security data', 'users/[id]/view, users/[id]'],
    ['GET', '/users/:id/profile', 'Profile data', 'users/[id]/view, users/[id]'],
    ['GET', '/users/:id/subscription', 'Subscription', 'users/[id]/view'],
    ['GET', '/users/:id/content', 'Content data', 'service only'],
    ['GET', '/users/:id/events', 'Events (paginated)', 'users/[id]/view'],
    ['GET', '/users/:id/community', 'Community (paginated)', 'users/[id]/view'],
    ['GET', '/users/:id/personal', 'Personal (paginated)', 'users/[id]/view, users/[id]'],
    ['GET', '/users/:id/activity', 'Activity (paginated)', 'users/[id]/view'],
    ['GET', '/sessions', 'Device sessions', '/admin/users/[id]'],
    ['GET', '/loginHistory', 'Login history', '/admin/users/[id]'],
    ['GET', '/posts', 'List posts', '/admin/community, posts list, post detail'],
    ['GET', '/posts/:id', 'Get post', '/admin/community/posts/[id]'],
    ['DELETE', '/posts/:id', 'Delete post', '/admin/community/posts'],
    ['GET', '/comments', 'List comments', '/admin/community, post detail, comments'],
    ['DELETE', '/comments/:id', 'Delete comment', '/admin/community/posts/[id], comments'],
    ['GET', '/reportedPosts', 'Reported posts', '/admin/community (counter), moderation, reported posts'],
    ['GET', '/reportedComments', 'Reported comments', '/admin/community (counter), moderation, reported comments'],
    ['PATCH', '/reportedPosts/:id', 'Update report status', 'ReportedList (reported posts)'],
    ['PATCH', '/reportedComments/:id', 'Update report status', 'ReportedList (reported comments)'],
    ['GET', '/safetyModeration', 'Moderation items', '/admin/community, moderation, safety/moderation'],
    ['GET', '/partners', 'Partners', '/admin/community/partners'],
    ['GET', '/anonymousPosts', 'Anonymous queue (paginated)', '/admin/community/anonymous'],
    ['GET', '/anonymousPosts/:postId', 'Anonymous post detail', '/admin/community/anonymous/[postId]/view'],
    ['POST', '/anonymousPosts/:postId/action', 'Moderation action', '/admin/community/anonymous/[postId]/view'],
    ['GET', '/anonymousModerationHistory', 'History (paginated)', '/admin/community/anonymous/history'],
    ['GET', '/dashboardStats', 'Dashboard stats', '/admin/dashboard'],
    ['GET', '/dashboardCharts', 'Dashboard chart data', '/admin/dashboard'],
    ['GET', '/dashboardTables', 'Dashboard recent activity', '/admin/dashboard'],
    ['GET', '/analyticsUsers', 'Analytics users', '/admin/analytics/users'],
    ['GET', '/payments', 'Payments', '/admin/finance/payments'],
    ['GET', '/revenue', 'Revenue', '/admin/finance/revenue'],
    ['GET', '/subscriptions', 'Subscriptions', '/admin/finance/subscriptions'],
    ['GET', '/support', 'Support tickets', '/admin/support'],
    ['GET', '/notifications', 'Notifications', '/admin/notifications'],
    ['GET', '/security', 'Security events', '/admin/security'],
    ['GET', '/auditLogs', 'Audit logs', '/admin/audit-logs'],
    ['GET', '/appSettings', 'App settings', '/admin/settings/app'],
    ['GET', '/trackers', 'Trackers', '/admin/settings/trackers'],
    ['GET', '/settings', 'Settings', '/admin/settings'],
    ['GET', '/aiDashboard', 'AI metrics', '/admin/ai/dashboard'],
    ['GET', '/aiUsage', 'AI usage', '/admin/ai/usage'],
    ['GET', '/aiConfig', 'AI config', '/admin/ai/config'],
    ['GET', '/safetyReports', 'Reports', '/admin/safety/reports'],
    ['GET', '/groups', 'Groups', '/admin/groups'],
    ['GET', '/events', 'Events', '/admin/events'],
    ['GET', '/trips', 'Trips', '/admin/trips'],
    ['GET', '/affirmations', 'Affirmations', '/admin/content/affirmations'],
    ['GET', '/banners', 'Banners', '/admin/content/banners'],
    ['GET', '/journalPrompts', 'Journal prompts', '/admin/content/journal'],
    ['GET', '/tips', 'Wellness tips', '/admin/content/tips'],
    ['GET', '/wardrobe', 'Wardrobe', '/admin/content/wardrobe'],
    ['GET', '/zodiac', 'Zodiac', '/admin/content/zodiac'],
])

# ============================================================
# 11. Page-to-API Mapping
# ============================================================
h1('11. Page-to-API Mapping')
table(['Page', 'API', 'Method', 'Purpose', 'Service'], [
    ['/login', '/auth/login', 'POST', 'Login', 'loginAdmin()'],
    ['useAuth() helper', '/auth/me', 'GET', 'Validate session', 'getCurrentAdmin()'],
    ['Sidebar logout', '/auth/logout', 'POST', 'Logout', 'logoutAdmin()'],
    ['/admin/dashboard', '/dashboardStats', 'GET', 'Role-keyed stats', 'direct'],
    ['/admin/dashboard', '/dashboardCharts', 'GET', 'Role-keyed chart', 'direct'],
    ['/admin/dashboard', '/dashboardTables', 'GET', 'Role-keyed activity table', 'direct'],
    ['/admin/users', '/users', 'GET', 'List users', 'getUsers()'],
    ['/admin/users', '/users/:id', 'PATCH', 'Inline edit save', 'updateUser()'],
    ['/admin/users', '/users/:id', 'DELETE', 'Delete user', 'deleteUser()'],
    ['/admin/users/[id]', '/users/:id', 'GET', 'Load user to edit', 'getUser()'],
    ['/admin/users/[id]', '/users/:id/profile', 'GET', 'Prepopulate profile', 'getUserProfile()'],
    ['/admin/users/[id]', '/users/:id/personal', 'GET', 'Prepopulate personal', 'getUserPersonal()'],
    ['/admin/users/[id]', '/users/:id', 'PATCH', 'Save edit', 'updateUser()'],
    ['/admin/users/[id]', '/sessions?userId=', 'GET', 'Device sessions', 'direct'],
    ['/admin/users/[id]', '/loginHistory?userId=', 'GET', 'Login history', 'direct'],
    ['/admin/users/[id]/view', '/users/:id', 'GET', 'Header info', 'getUser()'],
    ['/admin/users/[id]/view', '/users/:id/overview', 'GET', 'Overview tab', 'getUserOverview()'],
    ['/admin/users/[id]/view', '/users/:id/profile', 'GET', 'Profile tab', 'getUserProfile()'],
    ['/admin/users/[id]/view', '/users/:id/security', 'GET', 'Security tab', 'getUserSecurity()'],
    ['/admin/users/[id]/view', '/users/:id/subscription', 'GET', 'Subscription tab', 'getUserSubscription()'],
    ['/admin/users/[id]/view', '/users/:id/events', 'GET', 'Events tab', 'getUserEvents()'],
    ['/admin/users/[id]/view', '/users/:id/community', 'GET', 'Community tab', 'getUserCommunity()'],
    ['/admin/users/[id]/view', '/users/:id/personal', 'GET', 'Personal tab', 'getUserPersonal()'],
    ['/admin/users/[id]/view', '/users/:id/activity', 'GET', 'Activity tab', 'getUserActivity()'],
    ['/admin/users/[id]/view', '/users/:id', 'PATCH', 'Account status actions', 'updateUser()'],
    ['/admin/admins', '/admins', 'GET', 'List admins', 'getAdmins()'],
    ['/admin/admins', '/roles', 'GET', 'Resolve role names', 'getRoles()'],
    ['/admin/admins', '/admins/:id', 'DELETE', 'Delete admin', 'deleteAdmin()'],
    ['/admin/admins/create', '/admins', 'POST', 'Create admin', 'createAdmin()'],
    ['/admin/admins/[id]', '/admins/:id', 'GET', 'Load admin', 'getAdmin()'],
    ['/admin/admins/[id]', '/admins/:id', 'PATCH', 'Update admin', 'updateAdmin()'],
    ['/admin/roles', '/roles', 'GET', 'List roles', 'getRoles()'],
    ['/admin/roles', '/admins', 'GET', 'Admin counts', 'getAdmins()'],
    ['/admin/roles', '/roles/:id', 'DELETE', 'Delete role', 'deleteRole()'],
    ['/admin/roles/create', '/roles', 'POST', 'Create role', 'createRole()'],
    ['/admin/roles/[id]', '/roles/:id', 'GET', 'Load role', 'getRole()'],
    ['/admin/roles/[id]', '/roles/:id', 'PATCH', 'Update role', 'updateRole()'],
    ['/admin/community', '/posts', 'GET', 'Posts counter', 'direct'],
    ['/admin/community', '/comments', 'GET', 'Comments counter', 'direct'],
    ['/admin/community', '/reportedPosts', 'GET', 'Pending counter', 'direct'],
    ['/admin/community', '/reportedComments', 'GET', 'Pending counter', 'direct'],
    ['/admin/community', '/safetyModeration', 'GET', 'Pending counter', 'direct'],
    ['/admin/community', '/users', 'GET', 'Blocked users counter', 'direct'],
    ['/admin/community/moderation', '/reportedPosts', 'GET', 'Merged queue', 'direct'],
    ['/admin/community/moderation', '/reportedComments', 'GET', 'Merged queue', 'direct'],
    ['/admin/community/moderation', '/safetyModeration', 'GET', 'Merged queue', 'direct'],
])

table(['Page', 'API', 'Method', 'Purpose', 'Service'], [
    ['/admin/community/posts', '/posts', 'GET', 'List', 'direct'],
    ['/admin/community/posts', '/posts/:id', 'DELETE', 'Delete', 'direct'],
    ['/admin/community/posts/[id]', '/posts/:id', 'GET', 'Detail', 'direct'],
    ['/admin/community/posts/[id]', '/comments', 'GET', 'Comments for post', 'direct'],
    ['/admin/community/posts/[id]', '/comments/:id', 'DELETE', 'Delete comment', 'direct'],
    ['/admin/community/comments', '/comments', 'GET', 'List', 'direct'],
    ['/admin/community/comments', '/comments/:id', 'DELETE', 'Delete', 'direct'],
    ['/admin/community/reported/posts', '/reportedPosts', 'GET', 'List', 'direct (ReportedList)'],
    ['/admin/community/reported/posts', '/reportedPosts/:id', 'PATCH', 'Update status', 'direct (ReportedList)'],
    ['/admin/community/reported/comments', '/reportedComments', 'GET', 'List', 'direct (ReportedList)'],
    ['/admin/community/reported/comments', '/reportedComments/:id', 'PATCH', 'Update status', 'direct (ReportedList)'],
    ['/admin/community/blocked', '/users', 'GET', 'Filter blocked users', 'direct'],
    ['/admin/community/blocked', '/users/:id', 'PATCH', 'Unblock', 'direct'],
    ['/admin/community/anonymous', '/anonymousPosts', 'GET', 'Paginated queue (preserveEnvelope)', 'direct'],
    ['/admin/community/anonymous/[postId]/view', '/anonymousPosts/:postId', 'GET', 'Detail', 'direct'],
    ['/admin/community/anonymous/[postId]/view', '/anonymousPosts/:postId/action', 'POST', 'Moderate', 'direct'],
    ['/admin/community/anonymous/history', '/anonymousModerationHistory', 'GET', 'Paginated history', 'direct'],
    ['/admin/community/partners', '/partners', 'GET', 'List', 'direct'],
    ['/admin/groups', '/groups', 'GET', 'List', 'direct'],
    ['/admin/events', '/events', 'GET', 'List', 'direct'],
    ['/admin/trips', '/trips', 'GET', 'List', 'direct'],
    ['/admin/content/affirmations', '/affirmations', 'GET', 'List', 'direct'],
    ['/admin/content/banners', '/banners', 'GET', 'List', 'direct'],
    ['/admin/content/journal', '/journalPrompts', 'GET', 'List', 'direct'],
    ['/admin/content/tips', '/tips', 'GET', 'List', 'direct'],
    ['/admin/content/wardrobe', '/wardrobe', 'GET', 'List', 'direct'],
    ['/admin/content/zodiac', '/zodiac', 'GET', 'List', 'direct'],
    ['/admin/ai/dashboard', '/aiDashboard', 'GET', 'Metrics', 'direct'],
    ['/admin/ai/usage', '/aiUsage', 'GET', 'Usage', 'direct'],
    ['/admin/ai/config', '/aiConfig', 'GET', 'Config', 'direct'],
    ['/admin/analytics/users', '/analyticsUsers', 'GET', 'Analytics', 'direct'],
    ['/admin/support', '/support', 'GET', 'Tickets', 'direct'],
    ['/admin/notifications', '/notifications', 'GET', 'Notifications', 'direct'],
    ['/admin/security', '/security', 'GET', 'Events', 'direct'],
    ['/admin/audit-logs', '/auditLogs', 'GET', 'Audit logs', 'direct'],
    ['/admin/settings', '/settings', 'GET', 'Settings', 'direct'],
    ['/admin/settings/app', '/appSettings', 'GET', 'App settings', 'direct'],
    ['/admin/settings/trackers', '/trackers', 'GET', 'Trackers', 'direct'],
    ['/admin/finance/payments', '/payments', 'GET', 'Payments', 'direct'],
    ['/admin/finance/revenue', '/revenue', 'GET', 'Revenue', 'direct'],
    ['/admin/finance/subscriptions', '/subscriptions', 'GET', 'Subscriptions', 'direct'],
])

# ============================================================
# 12. API Data Flow
# ============================================================
h1('12. API Data Flow')
para('Example: viewing a user detail page.')
code('Browser admin\n'
     '  -> /admin/users/[id]/view (page.tsx)\n'
     '       -> useEffect -> getUser(userId)  [src/lib/api/users.ts]\n'
     '            -> fetchApi<User>("/users/<id>")  [src/lib/api/api.ts]\n'
     '                 -> fetch("http://localhost:3001/users/<id>",\n'
     '                      { headers: {"Authorization":"Bearer <token>"} })\n'
     '  <- { user fields... }  (normalized if paginated envelope)\n'
     '  -> setHeaderUser(user)\n'
     '  -> UI renders user cards/tables/tabs\n'
     'Error path: non-2xx -> fetchApi throws Error(body.message) -> caught in page ->\n'
     '   console.error + ErrorState / toast. 401 -> fetchApi clears session, hard redirect to /login.')
h2('12.1 Anonymous Queue Data Flow (paginated envelope)')
code('AnonymousPage.tsx\n'
     '  -> fetchApi<QueueResponse>("/anonymousPosts?filter=&search=&page=&limit=",\n'
     '       {}, preserveEnvelope=true)\n'
     '  -> response consumed AS-IS: response.data (array), response.page, response.pages\n'
     '  -> serverState for pagination; response is not normalized\n'
     '\n'
     'Reason: these responses use { data, total, pages, page } shape; preserveEnvelope keeps the\n'
     'whole envelope so the queue can read page/pages. fetchApi.normalizeResponse() would otherwise\n'
     'unwrap a {data,total,pages} envelope to just the data array (used for simple list pages).')

# ============================================================
# 13. State Management
# ============================================================
h1('13. State Management')
table(['Concern', 'Approach', 'Location'], [
    ['Global auth state', 'Module-level singleton (sharedAdmin, sharedLoading, authFetchPromise, subscribers set) shared via useAuth()', 'src/hooks/useAuth.ts'],
    ['Auth state writes', 'setAdmin / resetAuthState; logout clears and redirects', 'src/hooks/useAuth.ts'],
    ['Role/permission state', 'Module-level cache (cachedRoleId/cachedRole) + sessionStorage (role_<id>) via usePermissions()', 'src/hooks/usePermissions.ts'],
    ['Sidebar collapse state', 'React local state, persisted to localStorage key "sidebarExpanded"', 'src/hooks/useSidebar.ts'],
    ['Page/server data', 'Local useState per page component; useEffect fetch; no caching between navigations', 'each page'],
    ['UI local state', 'useState for search, filter, sort, pagination, edit/delete modals, toasts', 'each page'],
])
bullets([
    'No Redux, Zustand, React Context providers, or TanStack/React Query. No global cache invalidation across pages.',
    'Server state is not cached globally: navigating away and back to a page re-issues the same GET calls.',
    'Auth token/admin are stored in sessionStorage; role is cached in sessionStorage role_<roleId> by usePermissions().',
])

# ============================================================
# 14. Components and UI Architecture
# ============================================================
h1('14. Components and UI Architecture')
para('Reusable admin components live in src/components/admin/. Layout components (AdminLayout/Navbar/Sidebar/Footer) wrap every /admin route; page-level components are co-located under src/app/admin/<module>/.')
table(['Component', 'Purpose', 'Used by'], [
    ['AdminLayout', 'Shell: Navbar + Sidebar + Footer wrapper', 'src/app/admin/layout.tsx'],
    ['Sidebar', 'Permission-filtered navigation + logout button; collapse state persisted', 'AdminLayout'],
    ['Navbar', 'Top bar (title/breadcrumbs)', 'AdminLayout'],
    ['Footer', 'Footer', 'AdminLayout'],
    ['Loader', 'Spinner; used by login submit, loading states, and AdminLayout (full-screen)', 'multiple pages, login'],
    ['DashboardCardSkeleton', 'Loading placeholder for dashboard stat/chart/table cards', 'DashboardPage'],
    ['AdminTableSkeleton', 'Loading placeholder for table pages', 'users, admins, roles, community, etc.'],
    ['PermissionGate', 'Render children only when hasPermission(permission)', 'users, admins, roles pages, user view'],
    ['ConfirmModal', 'Yes/no confirmation dialog (delete, account actions)', 'users, admins, roles, user detail'],
    ['Toast', 'Timed success/error/info/warning toast', 'users, admins, roles, profile, user detail'],
    ['EmptyState / NoData', 'Empty / no-results display', 'community, moderation, history, tables'],
    ['ErrorState', 'Error display component present (not always wired)', 'shared/admin'],
    ['Skeleton', 'Skeleton loader (DashboardCardSkeleton exported)', 'components/admin/Skeleton'],
    ['PasswordInput', 'Password field component present', 'components/admin/PasswordInput'],
    ['Tabs', 'Tabs component used by user detail view', 'components/admin/Tabs'],
    ['Accordion', 'Accordion component present', 'components/admin/Accordion'],
    ['ReportedList', 'Shared reported-items list with status PATCH', 'community/reported/posts, reported/comments'],
])
para('Styling: Tailwind utility classes + CSS Modules per page (page.module.css). Icons from lucide-react. Charts via recharts LineChart on the dashboard.')

# ============================================================
# 15. Forms and Validation
# ============================================================
h1('15. Forms and Validation')
h2('15.1 Login form')
bullets([
    'Uses react-hook-form with zodResolver (@hookform/resolvers + zod) (src/app/login/page.tsx).',
    'Schema: email = z.string().email(), password = z.string().min(1).',
    'Validation messages: "Valid email is required", "Password is required". Inline display under each field.',
    'Password visibility toggle (Eye/EyeOff icon with onMouseDown preventDefault toggle).',
])
h2('15.2 Other forms')
bullets([
    'Users inline edit, Admins create/edit, Roles create/edit, and Profile use plain React state (useState) with no formal validation library — inputs are native and state is bound manually.',
    'Error/success feedback is shown via the Toast component (success/error) and inline error strings (e.g. "Failed to delete user").',
    'No Yup/other validation schema found for forms other than login.',
])

# ============================================================
# 16. Error Handling
# ============================================================
h1('16. Error Handling')
h2('16.1 API error handling')
bullets([
    'fetchApi() (src/lib/api/api.ts) is the single error boundary:',
    '  - 401 -> clearAuthToken(), sessionStorage.removeItem("admin"), window.location.href = "/login", throws "Session expired. Please log in again." (skipped on /login route).',
    '  - Other non-2xx -> parse body.message; fall back to "API error: <statusText>"; throw the Error.',
    '  - 204 -> returns undefined (no body).',
])
h2('16.2 UI error handling')
bullets([
    'Pages catch fetch rejections in .catch() and render an ErrorState component or a friendly message; some pages also console.error.',
    'UsersPage: setError("Failed to load users"). Admins/Roles: toast { type:"error" }.',
    'Login: setError(err.message || "Invalid email or password.").',
    'User detail view: caught errors set per-section error messages (e.g. getUserSecurity failure).',
])
note('There is no global error boundary provider; navigation to /admin without a session shows a full-screen Loader while loading, then null while redirecting.')

# ============================================================
# 17. Loading and Empty States
# ============================================================
h1('17. Loading and Empty States')
table(['Component / Location', 'Loading state', 'Empty / Error state'], [
    ['AdminLayout', 'full-screen Loader while auth loading; renders null when unauthenticated', 'redirects to /login'],
    ['DashboardPage', 'DashboardCardSkeleton placeholders for stat cards, chart and table', 'Activity table: "No activity data available." when no data'],
    ['UsersPage', 'AdminTableSkeleton', 'NoData / "Failed to load users" error'],
    ['AdminsPage / RolesPage', 'AdminTableSkeleton', 'NoData / error toast'],
    ['Community sub-pages', 'AdminTableSkeleton', 'NoData with title/description (e.g. "No moderation history")'],
    ['Anonymous queue', 'AdminTableSkeleton', 'NoData / error message'],
    ['Finance / Groups / Events / Trips / Analytics users / AI / Support / Security / Audit / Settings / Notifications', 'AdminTableSkeleton (where table present) or Loader', 'NoData when list empty'],
    ['Login', 'Loader spinner inside submit button', 'Inline error alert under form'],
    ['User detail tabs', 'Loader per section', '—']
])

# ============================================================
# 18. Environment Configuration
# ============================================================
h1('18. Environment Configuration')
table(['Variable', 'Purpose', 'Required', 'Sensitive'], [
    ['NEXT_PUBLIC_API_URL', 'Backend API base URL (used server-side as fallback when SSR-rendering the API base). In-browser the base is hardcoded to http://localhost:3001.', 'No (fallback to http://localhost:3001 used)', 'No (URL only, not a secret)'],
    ['MONGODB_URI', 'MongoDB connection string used by src/lib/mongo/mongoose.ts (server-side connect helper). Not imported by any app page.', 'No (fallback mongodb://127.0.0.1:27017/admin-dashboard exists)', 'Yes (connection string)'],
])
para('No .env files are committed at the repository root or in admin-dashboard-frontend. The client reads its API base from the hardcoded localhost:3001 in src/lib/api/api.ts and only consults NEXT_PUBLIC_API_URL on the server path. No secrets are exposed in the frontend code; the auth token is the only runtime secret and it is held in sessionStorage.')
note('Backend (admin-dashboard-backend) uses MONGODB_URI, PORT, TEST_MONGODB_URI, ENABLE_TEST_DB_SYNC — documented because frontend env resolution comments reference the API host; values are not present in the frontend repo.')

# ============================================================
# 19. Security Observations
# ============================================================
h1('19. Security Observations')
h2('19.1 Current Implementation (facts)')
bullets([
    'Bearer token is stored in sessionStorage under key "admin" (src/lib/api/api.ts). sessionStorage is tab-scoped and cleared on tab close, but is readable by any JavaScript running in the page (XSS risk).',
    'The same sessionStorage key "admin" also stores the full admin JSON, so login writes admin JSON and the token (TOKEN_KEY="admin") can overwrite each other depending on order (login writes JSON first, then token).',
    'Authorization header is sent on every fetch only when a token is present; login itself sends no token.',
    '401 responses clear the session and hard-redirect to /login.',
    'Permissions are enforced client-side only via PermissionGate and Sidebar filtering; there is no client-visible route guard beyond authentication.',
    'No dangerouslySetInnerHTML / innerHTML usage is present in the frontend (verified by search).',
    'No secrets are hardcoded; NEXT_PUBLIC_API_URL is the only build-time value and it is a URL, not a secret.',
])
h2('19.2 Potential Improvements (not present today)')
bullets([
    'Move the auth token to an httpOnly secure cookie to reduce XSS session-theft surface.',
    'Keep the token and the admin JSON in separate storage keys to avoid the overwrite observed during login.',
    'Add server-side route guards / middleware so permissions are enforced on the backend, not just hidden client-side (UI-level gating is not security).',
    'Avoid hard-coding http://localhost:3001 in getApiBaseUrl() (browser branch); rely purely on NEXT_PUBLIC_API_URL so non-localhost deployments work.',
    'Add retry / exponential-backoff and request cancellation (AbortController) for flaky networks.',
    'Surface server error details only to console and show generic toasts to avoid information leakage.',
])

# ============================================================
# 20. Known Gaps and TODOs
# ============================================================
h1('20. Known Gaps and TODOs')
bullets([
    'No source-level TODO/FIXME comments exist in src/ (only dependency integrity strings in package-lock.json matched the pattern). All identified gaps are inferred from inactive code.',
    'Users/[id]/view page contains a large fully commented-out implementation (lines ~1-980) above the active code — dead/commented code that should be removed.',
    'Analytics/features and Analytics/revenue pages are hardcoded mock stats with no API integration and no recharts charts.',
    'AI, Analytics, Finance, and Safety module index pages are static placeholders ("This is a placeholder page.").',
    'Content sub-page Journal is missing the /journalPrompts icon mapping and is listed as label "journal" instead of "Journal".',
    'getPermissions(), createUser(), replaceUser(), getUserContent(), registerAdmin() are defined in services but not used by any page.',
    'ProfilePage has no persistence: it builds form state from useAuth().admin but calls no API on save.',
    'Several Sidebar nav items and Community counters are commented out (anonymous/group counters; CMS sub-items; Settings/Profile footer links), indicating unfinished wiring.',
    'No automated tests in the frontend package.',
])

# ============================================================
# 21. Recommended Improvements
# ============================================================
h1('21. Recommended Improvements')
table(['Priority', 'Recommendation', 'Rationale'], [
    ['High', 'Back-end enforce permissions for every protected API', 'Client-side gating is UI only, not security'],
    ['High', 'Add server-side route guards / middleware redirect', 'Currently /admin relies on client useAuth; direct navigation can flash null content'],
    ['Medium', 'Replace hardcoded http://localhost:3001 with NEXT_PUBLIC_API_URL-only resolution', 'Enables non-localhost deployments'],
    ['Medium', 'Remove the commented-out users/[id]/view implementation and unused service functions', 'Dead code / tech debt'],
    ['Medium', 'Wire analytics/features, analytics/revenue to real APIs or mark clearly', 'Currently misleading hardcoded values'],
    ['Low', 'Implement missing POST/PUT/DELETE for content and other read-only modules if write features are desired', 'Currently read-only'],
    ['Low', 'Add unit/integration tests and CI', 'None present'],
])

# ============================================================
# 22. Developer Quick Start
# ============================================================
h1('22. Developer Quick Start')
h2('22.1 Prerequisites')
bullets(['Node.js 20+ (project built with Next 16 / React 19)', 'npm', 'MongoDB instance for the backend'])
h2('22.2 Setup')
code('# from admin-dashboard-frontend\n'
     'npm install\n\n'
     '# from admin-dashboard-backend\n'
     'npm install\n'
     'npm run seed   # optional; seeds the MongoDB API\n')
note('Commands below are run from the admin-dashboard-frontend project root; "npm" requires script-execution to be enabled or use "npm.cmd".')
h2('22.3 Commands')
table(['Goal', 'Command'], [
    ['Run frontend (dev)', 'npm run dev'],
    ['Build frontend', 'npm run build'],
    ['Run frontend (prod)', 'npm start'],
    ['Lint', 'npm run lint'],
    ['Run backend', 'npm run dev:api'],
    ['Seed backend', 'npm run seed'],
    ['Mock API (json-server, requires data/db.json)', 'npm run mock-api'],
])
note('The frontend dev server listens on http://localhost:3000 by default and calls the backend at http://localhost:3001. '
     'No .env is required for a local run; NEXT_PUBLIC_API_URL and MONGODB_URI have hardcoded fallbacks.')
h2('22.4 Key Files to Know')
table(['Purpose', 'File / Function'], [
    ['HTTP wrapper + token', 'src/lib/api/api.ts (fetchApi, getAuthToken, setAuthToken)'],
    ['Auth services', 'src/lib/api/auth.ts (loginAdmin, getCurrentAdmin, logoutAdmin, registerAdmin)'],
    ['User services', 'src/lib/api/users.ts (getUsers, getUser, createUser, updateUser, replaceUser, deleteUser, getUser* tabs)'],
    ['Admin services', 'src/lib/api/admins.ts (getAdmins, getAdmin, createAdmin, updateAdmin, deleteAdmin)'],
    ['Role services', 'src/lib/api/roles.ts (getRoles, getRole, createRole, updateRole, deleteRole, getPermissions)'],
    ['Auth state hook', 'src/hooks/useAuth.ts'],
    ['Permissions hook', 'src/hooks/usePermissions.ts'],
    ['Permission gate', 'src/components/admin/PermissionGate.tsx'],
    ['Navigation', 'src/components/admin/Sidebar/Sidebar.tsx'],
    ['Types', 'src/types/admin.ts, src/types/role.ts, src/lib/api/users.ts (User)'],
])

# ============================================================
# 23. Quick Reference
# ============================================================
h1('23. Quick Reference')
h2('23.1 Routes (admin section)')
bullets([
    '/admin/dashboard, /admin/users, /admin/users/[id], /admin/users/[id]/view,',
    '/admin/admins, /admin/admins/create, /admin/admins/[id],',
    '/admin/roles, /admin/roles/create, /admin/roles/[id],',
    '/admin/community, /admin/community/posts, /admin/community/posts/[id], /admin/community/comments,',
    '/admin/community/reported/posts, /admin/community/reported/comments, /admin/community/moderation, /admin/community/blocked,',
    '/admin/community/anonymous, /admin/community/anonymous/history, /admin/community/anonymous/[postId]/view, /admin/community/partners,',
    '/admin/groups, /admin/events, /admin/trips,',
    '/admin/content, /admin/content/affirmations, /admin/content/banners, /admin/content/journal, /admin/content/tips, /admin/content/wardrobe, /admin/content/zodiac,',
    '/admin/ai, /admin/ai/dashboard, /admin/ai/usage, /admin/ai/config,',
    '/admin/analytics, /admin/analytics/users, /admin/analytics/features, /admin/analytics/revenue,',
    '/admin/audit-logs, /admin/finance, /admin/finance/payments, /admin/finance/revenue, /admin/finance/subscriptions,',
    '/admin/notifications, /admin/profile, /admin/safety, /admin/safety/moderation, /admin/safety/reports,',
    '/admin/security, /admin/settings, /admin/settings/app, /admin/settings/trackers, /admin/support',
])
h2('23.2 Modules')
bullets(['Management: Users, Community (+anonymous moderation), Groups, Events, Trips, Content (CMS), Partners',
         'AI: AI Dashboard, AI Usage, AI Configuration',
         'Analytics: User Analytics, Feature Analytics, Revenue Analytics',
         'Finance: Subscriptions, Payments, Revenue',
         'Support: Support Tickets, Notifications',
         'Safety: Reports, Moderation Queue',
         'Administration: Admin Users, Roles & Permissions, Audit Logs, Security, Settings'])
h2('23.3 Authentication')
bullets(['Login: POST /auth/login -> Admin incl. token; token + admin JSON saved to sessionStorage key "admin".',
         'Session: GET /auth/me via shared useAuth() singleton (one call per load).',
         'Logout: POST /auth/logout then clear session (only on success).',
         '401 anywhere: clear session, hard redirect to /login.',
         'No refresh-token mechanism.'])
h2('23.4 State Management')
bullets(['useAuth() — module-level singleton (auth + redirect + logout).',
         'usePermissions() — module cache + sessionStorage role_<id>.',
         'useSidebar() — useState persisted in localStorage.',
         'Per-page useState for server data; no global server-state cache.',
         'No Redux/Zustand/Context providers/TanStack Query.'])
h2('23.5 Important Files')
bullets(['src/lib/api/api.ts — fetchApi wrapper, token storage, base URL, 401 handling',
         'src/lib/api/{auth,admins,roles,users}.ts — typed API service functions',
         'src/hooks/useAuth.ts, usePermissions.ts, useSidebar.ts',
         'src/components/admin/{PermissionGate,Toast,ConfirmModal,Loader,Skeleton,AdminLayout,Sidebar}.tsx',
         'src/types/admin.ts, src/types/role.ts'])

# ============================================================
# Diagrams (Mermaid)
# ============================================================
h1('Appendix A: Architecture Diagram (Mermaid)')
code('flowchart TD\n'
     '  AdminBrowser[Administrator browser]\n'
     '  NextApp["Next.js App Router (src/app)"]\n'
     '  Hooks["Custom hooks (useAuth/usePermissions/useSidebar)"\n'
     '  Components["UI + reusable components (src/components/admin)"]\n'
     '  APIServices["API services (src/lib/api: auth,admins,roles,users)"]\n'
     '  FetchApi[fetchApi() wrapper + token in sessionStorage]\n'
     '  Backend["Express API (localhost:3001)"]\n'
     '  Mongo["MongoDB (Mongoose)"]\n'
     '  AdminBrowser -->|HTTP| NextApp\n'
     '  NextApp -->|useAuth| Hooks\n'
     '  NextApp -->|render| Components\n'
     '  Components -->|service fn| APIServices\n'
     '  APIServices -->|fetch| FetchApi\n'
     '  FetchApi -->|Bearer token / JSON| Backend\n'
     '  Backend --> Mongo')
para('(Mermaid diagram above; paste into a Mermaid renderer to view.)')
h2('Authentication Flow (Mermaid)')
code('sequenceDiagram\n'
     '  participant U as Admin (browser)\n'
     '  participant L as LoginPage\n'
     '  participant F as fetchApi()\n'
     '  participant B as Backend\n'
     '  U->>L: submit email/password\n'
     '  L->>F: POST /auth/login {email,password}\n'
     '  F->>B: fetch /auth/login (no auth header)\n'
     '  B-->>F: Admin + token\n'
     '  F-->>L: Admin\n'
     '  L->>L: sessionStorage["admin"]=JSON, setAuthToken(token)\n'
     '  L->>U: router.push("/admin/dashboard")\n'
     '  U->>Next: navigate to /admin/*\n'
     '  Next->>F: GET /auth/me (Bearer token) [single shared call]\n'
     '  F->>B: /auth/me\n'
     '  B-->>F: Admin\n'
     '  F-->>Next: Admin\n'
     '  Next->>U: render AdminLayout (Navbar/Sidebar/Footer)\n'
     '  Note over F,B: Any 401 -> clear session -> redirect /login\n')
h2('API Request Flow (Mermaid)')
code('sequenceDiagram\n'
     '  participant P as Page (e.g. /admin/users)\n'
     '  participant H as useAuth / usePermissions\n'
     '  participant S as service fn (getUsers)\n'
     '  participant F as fetchApi()\n'
     '  participant B as Backend\n'
     '  P->>H: useAuth() / usePermissions()\n'
     '  H->>F: GET /auth/me ; GET /roles/:id\n'
     '  F->>B: /auth/me ; /roles/:id\n'
     '  P->>S: getUsers() on mount\n'
     '  S->>F: fetchApi("/users")\n'
     '  F->>B: GET /users (Bearer token)\n'
     '  B-->>F: users[]\n'
     '  F-->>S: users[]\n'
     '  S-->>P: users[]\n'
     '  P->>P: setUsers -> render table\n'
     '  P->>User: onClick delete\n'
     '  User->>S: deleteUser(id)\n'
     '  S->>F: DELETE /users/:id\n'
     '  F->>B: DELETE /users/:id\n')

# ============================================================
# Source References
# ============================================================
h2('Source Reference Index')
bullets([
    'HTTP wrapper: src/lib/api/api.ts (getApiBaseUrl, getAuthToken, setAuthToken, clearAuthToken, fetchApi, normalizeResponse)',
    'Auth services: src/lib/api/auth.ts (loginAdmin, getCurrentAdmin, logoutAdmin, registerAdmin)',
    'Admin services: src/lib/api/admins.ts (getAdmins, getAdmin, createAdmin, updateAdmin, deleteAdmin)',
    'Role services: src/lib/api/roles.ts (getRoles, getRole, createRole, updateRole, deleteRole, getPermissions)',
    'User services: src/lib/api/users.ts (getUsers, getUser, createUser, replaceUser, updateUser, deleteUser, getUser* tabs, ActivitySectionData, UserActivityResponse)',
    'Hooks: src/hooks/useAuth.ts, src/hooks/usePermissions.ts, src/hooks/useSidebar.ts',
    'Types: src/types/admin.ts (Admin), src/types/role.ts (Role, Permission), src/lib/api/users.ts (User)',
    'Components: src/components/admin/PermissionGate.tsx, Toast.tsx, ConfirmModal.tsx, Loader.tsx, Skeleton.tsx, ReportedList.tsx (under src/app/admin/community)',
    'Layout/nav: src/app/admin/layout.tsx, src/components/admin/AdminLayout/AdminLayout.tsx, Sidebar.tsx, Navbar.tsx, Footer.tsx',
    'Login: src/app/login/page.tsx (react-hook-form + zod schema in same file)',
    'Pages with API usage: dashboard, users, users/[id], users/[id]/view, admins, admins/create, admins/[id], roles, roles/create, roles/[id], community (+all sub-pages), groups, events, trips, content/*, ai/*, analytics/users, audit-logs, finance/*, notifications, safety/*, security, settings/*, support',
])

add_page_numbers()
h2('Appendix C: Validation Checklist')
bullets([
    'Routes considered: every /admin route from the file structure + /login + /403 + /not-found + root redirect.',
    'APIs considered: every fetchApi call captured via codebase grep across src/.',
    'Auth documented: login, /auth/me, token/session storage, 401 handling, logout, expiry.',
    'Permissions documented: PermissionGate, Sidebar filtering, per-component permission strings.',
    'Env documented: NEXT_PUBLIC_API_URL, MONGODB_URI; no secrets included.',
    'TODO/mock/placeholder identified: static placeholders, commented-out code, unused services.',
    'No secrets committed in this documentation.',
])

doc.save('d:/karthi/admin-dashboard/ADMIN_DASHBOARD_DOCUMENTATION.docx')
print('Saved ADMIN_DASHBOARD_DOCUMENTATION.docx')

# END-SENTINEL